"""The owner's lead spreadsheet: one clean row per venue, outreach first.

Built from the review queue plus contact details the official records
themselves publish (see Source.contact). No automatic outside lookups: the
Map, Google and Instagram columns are plain search links the owner clicks by
hand to find a phone number, website or social account.

The workbook (build_workbook) has a New tab, an All open tab, one tab per
state that has open leads, and a How scoring works tab. Every tab sorts by
lead score, highest first.

Local and email use only. Never print rows in GitHub Actions (public logs).
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from urllib.parse import quote_plus

from . import stage as stage_mod

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: (key, header, excel width). Order is the column order.
COLUMNS: list[tuple[str, str, int]] = [
    ("priority", "Priority", 9),
    ("hot", "Hot", 7),
    ("lead_score", "Score", 8),
    ("business_name", "Business name", 34),
    ("whats_new", "What's new", 16),
    ("company", "Company / owner", 34),
    ("business_type", "Business type", 26),
    ("filing", "Filing", 18),
    ("stage", "Stage", 11),
    ("filed_on", "Filed on", 12),
    ("phone", "Phone", 16),
    ("people", "Owner / applicant names", 34),
    ("address", "Address", 34),
    ("city", "City", 18),
    ("state", "State", 7),
    ("zip", "ZIP", 8),
    ("market", "Market", 24),
    ("mailing_address", "Mailing address", 40),
    ("license", "License applied for", 40),
    ("map_url", "Map", 8),
    ("google_url", "Google", 9),
    ("instagram_url", "Instagram", 11),
    ("record_url", "Official record", 14),
    ("lead_ids", "Lead ID", 12),
]
HEADERS = [h for _, h, _ in COLUMNS]
#: All open and state tabs span many days: the queue date replaces What's new.
OPEN_COLUMNS = [("queue_date", "Queued on", 12) if key == "whats_new" else (key, h, w)
                for key, h, w in COLUMNS]

NEW_FILING = "New filing"
STAGE_ADVANCED = "Stage advanced"
DETAILS_CHANGED = "Details changed"

# Tier A/B mean nightclub / bar-or-venue (see qualify.py); C is by license.
BUSINESS_TYPES = {
    "nightlife": "Restaurant / bar",
    "on_premise": "Restaurant",
    "hospitality_mfg": "Brewery / taproom / winery",
    "catering_event": "Caterer / event venue",
    "hotel": "Hotel",
}
_TYPE_ORDER = list(BUSINESS_TYPES)

# (pattern on the upper-cased application type, plain label). First match wins.
_FILINGS = [
    (r"NEW LICENSE", "Newly licensed"),
    (r"\bNEW\b|ORIGINAL", "New application"),
    (r"ASSUMPTION|TRANSFER|CHANGE OF (OWNER|CORP|OFFICER|STOCK)|OWNERSHIP", "Change of owner"),
    (r"LOCATION|RELOC|C_LOC", "New location"),
    (r"TRADENAME|TRADE NAME|NAME CHANGE", "Name change"),
    (r"CLASS|IN LIEU|UPGRADE|ADDED PRIVILEGE", "License upgrade"),
]


def _business_type(priority: str, cats: list) -> str:
    if priority == "A":
        return "Nightclub / lounge"
    if priority == "B":
        return ("Brewery / taproom" if cats and cats[0] == "hospitality_mfg"
                else "Bar / event venue")
    return BUSINESS_TYPES.get(cats[0], "Other") if cats else "Other"


def _search(base: str, query: str) -> str | None:
    """A plain search link. Instagram's own search needs a login, so the
    Instagram column is a Google search limited to instagram.com."""
    query = " ".join((query or "").split())
    return base + quote_plus(query) if query else None


def _filing(app_types: list[str], source: str) -> str:
    labels = []
    for t in app_types:
        up = t.upper()
        label = next((lab for pat, lab in _FILINGS if re.search(pat, up)), t.title())
        if label not in labels:
            labels.append(label)
    if not labels and source == "ca_abc_applications":
        return "New application"  # CA export lists pending applications only
    return ", ".join(labels)


def _stage(recs: list[dict], sources: dict) -> str:
    """Most advanced stage across a venue's records. Uses the stored stage,
    or the source's mapping for rows stored before stages existed."""
    found = []
    for rec in recs:
        value = rec.get("stage")
        if value is None:
            src = sources.get(rec.get("source"))
            value = (_source_stage(src, rec) if src
                     else stage_mod.from_status(rec.get("status")))
        found.append(value)
    return stage_mod.best(found) or ""


def _source_stage(src, rec: dict) -> str | None:
    from .models import Record

    try:
        return src.stage(Record(source=rec.get("source") or "", source_record_id="",
                                source_url="", status=rec.get("status"),
                                license_type=rec.get("license_type")))
    except Exception:  # noqa: BLE001 - a label only
        return None


def _whats_new(recs: list[dict]) -> str:
    """New filing > Stage advanced > Details changed, from the queue events."""
    types = {r.get("event_type") for r in recs}
    if types & {"new", "baseline"}:
        return NEW_FILING
    if any(r.get("event_type") == "changed" and "stage" in (r.get("changes") or {})
           for r in recs):
        return STAGE_ADVANCED
    return DETAILS_CHANGED if "changed" in types else ""


def _phone(value: str | None) -> str | None:
    digits = re.sub(r"\D", "", value or "")
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) == 10:
        return f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    return (value or "").strip() or None


def _title(text: str | None) -> str:
    """Tidy ALL-CAPS names; leave mixed case as the source wrote it."""
    text = " ".join((text or "").split())
    if text and text == text.upper() and any(c.isalpha() for c in text):
        text = text.title()
        text = re.sub(r"(\d)(St|Nd|Rd|Th)\b", lambda m: m.group(1) + m.group(2).lower(), text)
        text = re.sub(r"'S\b", "'s", text)
        text = re.sub(r"\b([A-Z][a-z])(?= \d{5})", lambda m: m.group(1).upper(), text)
    return re.sub(r"\b(Llc|Inc|Llp|Pllc|Ltd|Dba|Usa)\b", lambda m: m.group(1).upper(), text)


def _people(contacts: list[dict], names: tuple) -> str | None:
    """Owner/applicant names, minus entries that just repeat the business."""
    skip = {re.sub(r"\W", "", (n or "").upper()) for n in names if n}
    people = []
    for c in contacts:
        for p in (c.get("people") or "").split(";"):
            p = _title(p.strip())
            if p and re.sub(r"\W", "", p.upper()) not in skip and p not in people:
                people.append(p)
    return "; ".join(people) or None


def _uniq(values) -> list[str]:
    out: list[str] = []
    for v in values:
        v = (v or "").strip() if isinstance(v, str) else v
        if v and v not in out:
            out.append(v)
    return out


_SQL = """
SELECT q.queue_date, q.record_id, q.venue_key, q.tier, q.score, q.legal_name,
       q.dba, q.license_description, q.application_type, q.status,
       q.application_date, q.address, q.city, q.state, q.zip, q.metro,
       q.source, q.source_url, r.category, r.raw, q.event_type, q.changes,
       q.review_status, q.stage, q.lead_score, q.hot, q.license_type,
       q.first_seen_at
FROM review_queue q JOIN records r ON r.id = q.record_id
WHERE r.qualified
  AND (%(day)s::date IS NULL OR q.queue_date = %(day)s)
  AND (NOT %(open)s OR q.review_status = 'new')
ORDER BY q.queue_date DESC, q.score DESC, q.record_id
"""
_FIELDS = ["queue_date", "record_id", "venue_key", "tier", "score", "legal_name", "dba",
           "license_description", "application_type", "status", "application_date",
           "address", "city", "state", "zip", "metro", "source", "source_url",
           "category", "raw", "event_type", "changes", "review_status", "stage",
           "lead_score", "hot", "license_type", "first_seen_at"]


def load_rows(conn, day: date | None, open_only: bool = False) -> list[dict]:
    """Queued leads for `day` (all days if None), one row per venue per day."""
    with conn.cursor() as cur:
        cur.execute(_SQL, {"day": day, "open": open_only})
        records = [dict(zip(_FIELDS, r)) for r in cur.fetchall()]
    return group_records(records)


def group_records(records: list[dict]) -> list[dict]:
    """Merge application records (best score first) into venue rows."""
    from .sources import all_sources

    by_source = {s.name: s for s in all_sources()}
    groups: dict[tuple, list[dict]] = {}
    for rec in records:
        groups.setdefault((rec["queue_date"], rec["venue_key"]), []).append(rec)

    rows = []
    for (queue_date, _), recs in groups.items():
        recs.sort(key=lambda r: (-(r.get("score") or 0), r["record_id"]))
        top = recs[0]
        contacts = []
        for rec in recs:
            src = by_source.get(rec["source"])
            try:
                contacts.append(src.contact(rec.get("raw") or {}) if src else {})
            except Exception:  # noqa: BLE001 - contact details are best effort
                contacts.append({})
        dates = [r["application_date"] for r in recs if r.get("application_date")]
        cats = sorted({r.get("category") for r in recs},
                      key=lambda c: _TYPE_ORDER.index(c) if c in _TYPE_ORDER else 99)
        name = top.get("dba") or top.get("legal_name")
        company = top.get("legal_name") if top.get("legal_name") != name else None
        address = ", ".join(p for p in (top.get("address"), top.get("city"),
                                        top.get("state"), top.get("zip")) if p)
        map_q = " ".join(p for p in (name, address) if p)
        web_q = " ".join(p for p in (name, top.get("city"), top.get("state")) if p)
        seen = [r["first_seen_at"] for r in recs if r.get("first_seen_at")]
        rows.append({
            "queue_date": queue_date,
            "venue_key": top.get("venue_key"),
            "priority": min((r.get("tier") or "C") for r in recs),
            "score": top.get("score") or 0,
            "lead_score": max((r.get("lead_score") or 0) for r in recs),
            "hot": "Hot" if any(r.get("hot") for r in recs) else "",
            "whats_new": _whats_new(recs),
            "first_seen": min(seen) if seen else None,
            "business_name": _title(name),
            "company": _title(company) or None,
            "business_type": _business_type(min((r.get("tier") or "C") for r in recs), cats),
            "filing": _filing(_uniq(r.get("application_type") for r in recs), top["source"]),
            "stage": _stage(recs, by_source),
            "filed_on": min(dates) if dates else None,
            "phone": ", ".join(_uniq(_phone(c.get("phone")) for c in contacts)) or None,
            "people": _people(contacts, (name, top.get("legal_name"))),
            "address": _title(top.get("address")),
            "city": _title(top.get("city")),
            "state": (top.get("state") or "").upper(),
            "zip": top.get("zip") or "",
            "market": top.get("metro") or "",
            "mailing_address": _title(next((c["mailing_address"] for c in contacts
                                            if c.get("mailing_address")), None)) or None,
            "license": "; ".join(_uniq(r.get("license_description") for r in recs)),
            "map_url": _search("https://www.google.com/maps/search/?api=1&query=", map_q),
            "google_url": _search("https://www.google.com/search?q=", web_q),
            "instagram_url": _search("https://www.google.com/search?q=",
                                     f"site:instagram.com {name}" if name else ""),
            "record_url": top.get("source_url"),
            "lead_ids": " ".join(str(r["record_id"]) for r in
                                 sorted(recs, key=lambda r: r["record_id"])),
        })
    rows.sort(key=sort_key)
    return rows


def sort_key(row: dict):
    """Highest lead score first, then priority, market, name."""
    return (-(row.get("lead_score") or 0), row.get("priority") or "C",
            row.get("market") or "", row.get("business_name") or "")


def _cell(value):
    if isinstance(value, datetime):
        return value.date()
    return "" if value is None else value


def write_csv(rows: list[dict], out) -> None:
    w = csv.writer(out)
    w.writerow(HEADERS)
    for row in rows:
        w.writerow([_csv_safe(_cell(row.get(k))) for k, _, _ in COLUMNS])


def _csv_safe(value):
    """Stop spreadsheet apps from running a source string as a formula."""
    if isinstance(value, str) and value[:1] in ("=", "+", "@"):
        return "'" + value
    return value


def build_xlsx(rows: list[dict], title: str = "Leads") -> bytes:
    """One sheet of venue rows (the New tab layout)."""
    from openpyxl import Workbook

    wb = Workbook()
    _write_sheet(wb.active, rows, COLUMNS, title)
    return _save(wb)


def build_workbook(new_rows: list[dict], open_rows: list[dict] | None = None) -> bytes:
    """The owner's workbook: New, All open, one tab per state with open
    leads (from the data, not a fixed list), then How scoring works."""
    from openpyxl import Workbook

    open_rows = sorted(open_rows or [], key=sort_key)
    wb = Workbook()
    _write_sheet(wb.active, sorted(new_rows, key=sort_key), COLUMNS, "New")
    _write_sheet(wb.create_sheet(), open_rows, OPEN_COLUMNS, "All open")
    for state in sorted({r.get("state") or "Other" for r in open_rows}):
        _write_sheet(wb.create_sheet(),
                     [r for r in open_rows if (r.get("state") or "Other") == state],
                     OPEN_COLUMNS, state)
    _write_legend(wb.create_sheet())
    return _save(wb)


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_sheet(ws, rows: list[dict], columns: list[tuple[str, str, int]],
                 title: str) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws.title = re.sub(r"[\[\]:*?/\\]", "-", title)[:31] or "Leads"
    ws.append([h for _, h, _ in columns])
    head_fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22
    fills = {"A": PatternFill("solid", fgColor="D1FAE5"),
             "B": PatternFill("solid", fgColor="FEF3C7"),
             "C": PatternFill("solid", fgColor="F3F4F6")}
    hot_fill = PatternFill("solid", fgColor="FEE2E2")
    link_font = Font(color="1D4ED8", underline="single")
    link_cols = {"map_url": "Map", "google_url": "Search", "instagram_url": "Search",
                 "record_url": "Record"}
    keys = [key for key, _, _ in columns]
    for i, row in enumerate(rows, start=2):
        for j, key in enumerate(keys, start=1):
            value = _cell(row.get(key))
            cell = ws.cell(row=i, column=j)
            if key in link_cols:
                if isinstance(value, str) and value.startswith(("http://", "https://")):
                    cell.value = link_cols[key]
                    cell.hyperlink = value
                    cell.font = link_font
                continue
            cell.value = value
            if cell.data_type == "f":
                cell.data_type = "s"  # never let a source string become a formula
            if isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
        prio = ws.cell(row=i, column=keys.index("priority") + 1)
        prio.alignment = Alignment(horizontal="center")
        if row.get("priority") in fills:
            prio.fill = fills[row["priority"]]
            prio.font = Font(bold=True)
        if row.get("hot"):
            hot = ws.cell(row=i, column=keys.index("hot") + 1)
            hot.fill = hot_fill
            hot.font = Font(bold=True, color="B91C1C")
    for j, (_, _, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = f"{get_column_letter(keys.index('business_name') + 2)}2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"


def legend_lines() -> list[tuple[str, str]]:
    """(label, value) rows for the How scoring works tab, built from the
    live points table in qualify.py so the two never drift."""
    from . import qualify as q

    lines = [
        ("How scoring works", ""),
        ("", ""),
        ("Score", "0 to 100. Higher means a better ticketing fit. Parts add up."),
        ("Hot", f"A nightclub or lounge (priority A) with a score of "
                f"{q.hot_min_score()} or more."),
        ("", ""),
        ("Venue type", "Points"),
        ("A: nightclub or lounge", q.TIER_POINTS["A"]),
        ("B: bar or event venue", q.TIER_POINTS["B"]),
        ("C: restaurant", q.TIER_POINTS["C"]),
        ("", ""),
        ("Nightlife license (highest one counts)", "Points"),
    ]
    lines += [(label, pts) for pts, label in sorted(q.NIGHTLIFE_LICENSE_POINTS.values(),
                                                   key=lambda v: -v[0])]
    lines += [("", ""), ("Stage", "Points")]
    lines += [(s, q.STAGE_POINTS[s]) for s in stage_mod.STAGES]
    lines += [
        ("", ""),
        ("Filing type", "Points"),
        ("New filing or new location", q.FILING_POINTS[0][1]),
        ("Change of owner", next(p for _, p, lab in q.FILING_POINTS
                                 if lab == "change of owner")),
        ("", ""),
        ("Stages", ""),
        ("Licensed", "Issued or active. Florida counts only if issued in the last 60 days."),
        ("Approved", "Approved or conditional, not yet active."),
        ("In review", "Past intake, in process."),
        ("Received", "Just filed."),
        ("", ""),
        ("What's new", ""),
        (NEW_FILING, "First time this filing showed up."),
        (STAGE_ADVANCED, "A filing we already had moved to a later stage."),
        (DETAILS_CHANGED, "Name, address, license or status changed."),
        ("", ""),
        ("Tabs", ""),
        ("New", "The day's leads."),
        ("All open", "Every lead not yet reviewed, from all days."),
        ("State tabs", "All open, split by state."),
    ]
    return lines


def _write_legend(ws) -> None:
    from openpyxl.styles import Font

    ws.title = "How scoring works"
    for label, value in legend_lines():
        ws.append([label, value])
        if value == "Points" or label in ("How scoring works", "Stages", "What's new",
                                           "Tabs"):
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 80
