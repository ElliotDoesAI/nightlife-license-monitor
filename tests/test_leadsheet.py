"""The owner's lead spreadsheet. Synthetic data only."""

import csv
import io
from datetime import date

from openpyxl import load_workbook

from licmon import leadsheet
from licmon.sources.ca_abc import CaAbcSource
from licmon.sources.fl_abt import FlAbtSource
from licmon.sources.wa_lcb import WaLcbSource

DAY = date(2026, 10, 1)


def rec(record_id, **kw):
    base = {
        "queue_date": DAY, "record_id": record_id, "venue_key": f"k{record_id}",
        "tier": "B", "score": 60, "legal_name": None, "dba": None,
        "license_description": "Mixed Beverage Permit", "application_type": "ORIGINAL",
        "status": "Received", "application_date": date(2026, 9, 20),
        "address": "1 FAKE ST", "city": "AUSTIN", "state": "TX", "zip": "78701",
        "metro": "Austin", "source": "tx_tabc_pending",
        "source_url": "https://example.invalid/r", "category": "on_premise", "raw": {},
    }
    base.update(kw)
    return base


def test_contact_hooks_use_only_the_official_row():
    wa = WaLcbSource().contact({"Contact Phone": "2065550100",
                                "Applicant(s)": "ZEBRA FAKE LLC; JANE Q TESTER"})
    assert wa == {"phone": "2065550100", "people": "ZEBRA FAKE LLC; JANE Q TESTER"}
    ca = CaAbcSource().contact({"rows": [
        {"Mail Addr 1": "", "Mail City": ""},
        {"Mail Addr 1": "PO BOX 1", "Mail Addr 2": "", "Mail City": "FAKEVILLE",
         "Mail State": "CA", "Mail Zip": "90000"}]})
    assert ca == {"mailing_address": "PO BOX 1, FAKEVILLE, CA 90000"}
    fl = FlAbtSource().contact({"Mail Address 1": "9 TEST WAY", "Mail Address 2": "STE 2",
                                "Mail City": "MIAMI", "Mail State": "FL", "Mail ZIP": "33101"})
    assert fl == {"mailing_address": "9 TEST WAY STE 2, MIAMI, FL 33101"}
    assert FlAbtSource().contact({}) == {}


def test_group_merges_venue_and_cleans_fields():
    rows = leadsheet.group_records([
        rec(1, venue_key="v", tier="A", score=80, legal_name="ZEBRA FAKE LLC",
            dba="ZEBRA FAKE LOUNGE", license_description="Late Hours Certificate",
            category="nightlife", application_date=date(2026, 9, 18)),
        rec(2, venue_key="v", tier="B", score=60, legal_name="ZEBRA FAKE LLC",
            dba="ZEBRA FAKE LOUNGE", status="Pending – In Review"),
        rec(3, legal_name="Other Fake Inc", dba=None, tier="C", score=45),
    ])
    assert [r["business_name"] for r in rows] == ["Zebra Fake Lounge", "Other Fake INC"]
    top = rows[0]
    assert top["priority"] == "A"
    assert top["company"] == "Zebra Fake LLC"
    assert top["business_type"] == "Nightclub / lounge"
    assert top["filing"] == "New application"
    assert top["stage"] == "In review"
    assert top["filed_on"] == date(2026, 9, 18)
    assert top["license"] == "Late Hours Certificate; Mixed Beverage Permit"
    assert top["lead_ids"] == "1 2"
    assert top["map_url"].startswith("https://www.google.com/maps/search/?api=1&query=ZEBRA")
    assert top["google_url"] == "https://www.google.com/search?q=ZEBRA+FAKE+LOUNGE+AUSTIN+TX"
    assert top["instagram_url"] == ("https://www.google.com/search?q="
                                    "site%3Ainstagram.com+ZEBRA+FAKE+LOUNGE")
    assert rows[1]["company"] is None  # no DBA: the name is the company


def test_wa_phone_and_people_without_repeating_business():
    [row] = leadsheet.group_records([rec(
        7, source="wa_lcb_actions", dba="ZEBRA FAKE BAR", legal_name=None,
        application_type="ASSUMPTION", status="APPROVED",
        raw={"Contact Phone": "206-555-0100",
             "Applicant(s)": "ZEBRA FAKE BAR; JANE Q TESTER; JOHN TESTER"})])
    assert row["phone"] == "(206) 555-0100"
    assert row["people"] == "Jane Q Tester; John Tester"
    assert row["filing"] == "Change of owner"
    assert row["stage"] == "Licensed"  # WA "approved" section = issued


def test_filing_and_status_labels():
    assert leadsheet._filing(["NEW LICENSE"], "fl_abt_licenses") == "Newly licensed"
    assert leadsheet._filing([], "ca_abc_applications") == "New application"
    assert leadsheet._filing(["ADDED/CHANGE OF TRADENAME"], "x") == "Name change"
    by_source = {"fl": FlAbtSource(), "ca": CaAbcSource()}

    def stage_of(status, source):
        return leadsheet._stage([{"status": status, "source": source}], by_source)

    assert stage_of("Current", "fl") == "Licensed"
    assert stage_of("PEND", "ca") == "Received"
    assert stage_of("IntakeComplete", "unknown") == "Received"
    assert stage_of("Withdrawn (application withdrawn)", "fl") == ""
    # a stored stage wins over the status text
    assert leadsheet._stage([{"status": "PEND", "stage": "Approved", "source": "ca"}],
                            by_source) == "Approved"


def test_xlsx_layout_links_and_no_formulas():
    rows = leadsheet.group_records([
        rec(1, tier="A", score=80, dba="=HYPERLINK(\"http://evil.invalid\")"),
        rec(2, dba="ZEBRA FAKE TAVERN")])
    wb = load_workbook(io.BytesIO(leadsheet.build_xlsx(rows, title="Leads 2026-10-01")))
    ws = wb.active
    assert ws.title == "Leads 2026-10-01"
    assert [c.value for c in ws[1]] == leadsheet.HEADERS
    assert ws.freeze_panes == "E2"  # through Business name
    assert ws.auto_filter.ref == "A1:X3"
    name_col = leadsheet.HEADERS.index("Business name") + 1
    evil = ws.cell(row=2, column=name_col)
    assert evil.data_type == "s" and evil.value.startswith("=")
    rec_col = leadsheet.HEADERS.index("Official record") + 1
    assert ws.cell(row=2, column=rec_col).value == "Record"
    assert ws.cell(row=2, column=rec_col).hyperlink.target == "https://example.invalid/r"
    date_col = leadsheet.HEADERS.index("Filed on") + 1
    assert ws.cell(row=2, column=date_col).value.date() == date(2026, 9, 20)


def test_csv_has_headers_and_neutralizes_formulas():
    rows = leadsheet.group_records([rec(1, dba="=1+1")])
    buf = io.StringIO()
    leadsheet.write_csv(rows, buf)
    out = list(csv.reader(io.StringIO(buf.getvalue())))
    assert out[0] == leadsheet.HEADERS
    assert out[1][leadsheet.HEADERS.index("Business name")] == "'=1+1"


def test_empty_sheet_is_valid():
    wb = load_workbook(io.BytesIO(leadsheet.build_xlsx([])))
    assert [c.value for c in wb.active[1]] == leadsheet.HEADERS


def test_score_hot_and_whats_new_per_venue():
    rows = leadsheet.group_records([
        rec(1, venue_key="v", tier="A", lead_score=85, hot=True, stage="Approved",
            event_type="changed", changes={"stage": ["Received", "Approved"]}),
        rec(2, venue_key="v", tier="A", lead_score=60, hot=False, stage="Received",
            event_type="changed", changes={"status": ["A", "B"]}),
        rec(3, venue_key="w", tier="B", lead_score=40, event_type="changed",
            changes={"address": ["1 A ST", "2 A ST"]}),
        rec(4, venue_key="x", tier="C", lead_score=90, event_type="new"),
    ])
    by_key = {r["venue_key"]: r for r in rows}
    assert by_key["v"]["lead_score"] == 85 and by_key["v"]["hot"] == "Hot"
    assert by_key["v"]["stage"] == "Approved"
    assert by_key["v"]["whats_new"] == leadsheet.STAGE_ADVANCED
    assert by_key["w"]["whats_new"] == leadsheet.DETAILS_CHANGED
    assert by_key["x"]["whats_new"] == leadsheet.NEW_FILING
    # sorted by score, highest first, whatever the tier
    assert [r["venue_key"] for r in rows] == ["x", "v", "w"]


def test_workbook_tabs_come_from_the_data():
    new = leadsheet.group_records([
        rec(1, tier="A", lead_score=80, hot=True, dba="ZEBRA FAKE LOUNGE", event_type="new"),
        rec(2, tier="B", lead_score=50, dba="ZEBRA FAKE TAVERN", event_type="new")])
    open_rows = leadsheet.group_records([
        rec(1, tier="A", lead_score=80, hot=True, dba="ZEBRA FAKE LOUNGE"),
        rec(5, tier="C", lead_score=30, dba="FAKE BISTRO", state="CA", city="LA",
            queue_date=date(2026, 9, 28)),
        rec(6, tier="B", lead_score=95, dba="FAKE TAPROOM", state="WA",
            queue_date=date(2026, 9, 27))])
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook(new, open_rows)))
    assert wb.sheetnames == ["New", "All open", "CA", "TX", "WA", "How scoring works"]
    new_ws, open_ws = wb["New"], wb["All open"]
    assert [c.value for c in new_ws[1]] == leadsheet.HEADERS
    assert "What's new" in leadsheet.HEADERS and "Stage" in leadsheet.HEADERS
    open_headers = [c.value for c in open_ws[1]]
    assert "Queued on" in open_headers and "What's new" not in open_headers
    hot_col = leadsheet.HEADERS.index("Hot") + 1
    score_col = leadsheet.HEADERS.index("Score") + 1
    assert new_ws.cell(row=2, column=hot_col).value == "Hot"
    assert [new_ws.cell(row=r, column=score_col).value for r in (2, 3)] == [80, 50]
    assert new_ws.cell(row=2, column=leadsheet.HEADERS.index("What's new") + 1).value \
        == leadsheet.NEW_FILING
    # All open: every day, highest score first
    names = [open_ws.cell(row=r, column=open_headers.index("Business name") + 1).value
             for r in (2, 3, 4)]
    assert names == ["Fake Taproom", "Zebra Fake Lounge", "Fake Bistro"]
    assert wb["CA"].max_row == 2 and wb["TX"].max_row == 2
    legend = [c.value for c in wb["How scoring works"]["A"]]
    assert "Public place of amusement (Chicago)" in legend and "Licensed" in legend
    text = " ".join(str(v) for row in wb["How scoring works"].values for v in row if v)
    assert "—" not in text  # no em dashes in owner-facing text


def test_workbook_without_open_rows_is_valid():
    wb = load_workbook(io.BytesIO(leadsheet.build_workbook([])))
    assert wb.sheetnames == ["New", "All open", "How scoring works"]
